import re
from dataclasses import dataclass, field
from pathlib import Path

from app.core.errors import AppError


@dataclass(slots=True)
class ParsedBlock:
    text: str
    page_no: int | None = None
    slide_no: int | None = None
    section_title: str | None = None
    metadata: dict = field(default_factory=dict)


@dataclass(slots=True)
class ParsedDocument:
    blocks: list[ParsedBlock]
    metadata: dict = field(default_factory=dict)


class BaseParser:
    def parse(self, file_path: Path) -> ParsedDocument:
        raise NotImplementedError


class PdfParser(BaseParser):
    def parse(self, file_path: Path) -> ParsedDocument:
        try:
            import fitz
        except ImportError as exc:
            raise AppError(
                "PARSE_DEPENDENCY_MISSING",
                "PyMuPDF is required to parse PDF files.",
                details={"package": "pymupdf"},
            ) from exc

        blocks: list[ParsedBlock] = []
        try:
            with fitz.open(file_path) as document:
                for page_index, page in enumerate(document, start=1):
                    text = page.get_text("text")
                    if text.strip():
                        blocks.append(ParsedBlock(text=text, page_no=page_index))
        except Exception as exc:
            raise AppError("PARSE_FAILED", "PDF parsing failed.") from exc
        return _ensure_non_empty(ParsedDocument(blocks=blocks, metadata={"parser": "pdf"}))


class DocxParser(BaseParser):
    def parse(self, file_path: Path) -> ParsedDocument:
        try:
            from docx import Document
        except ImportError as exc:
            raise AppError(
                "PARSE_DEPENDENCY_MISSING",
                "python-docx is required to parse DOCX files.",
                details={"package": "python-docx"},
            ) from exc

        blocks: list[ParsedBlock] = []
        current_title: str | None = None
        try:
            document = Document(file_path)
            for paragraph in document.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                style_name = paragraph.style.name if paragraph.style else ""
                if style_name.lower().startswith("heading"):
                    current_title = text[:255]
                blocks.append(
                    ParsedBlock(
                        text=text,
                        section_title=current_title,
                        metadata={"style": style_name},
                    )
                )
        except Exception as exc:
            raise AppError("PARSE_FAILED", "DOCX parsing failed.") from exc
        return _ensure_non_empty(ParsedDocument(blocks=blocks, metadata={"parser": "docx"}))


class PptxParser(BaseParser):
    def parse(self, file_path: Path) -> ParsedDocument:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise AppError(
                "PARSE_DEPENDENCY_MISSING",
                "python-pptx is required to parse PPTX files.",
                details={"package": "python-pptx"},
            ) from exc

        blocks: list[ParsedBlock] = []
        try:
            presentation = Presentation(file_path)
            for slide_index, slide in enumerate(presentation.slides, start=1):
                shape_text = [
                    shape.text.strip()
                    for shape in slide.shapes
                    if hasattr(shape, "text") and shape.text and shape.text.strip()
                ]
                if shape_text:
                    title = shape_text[0][:255]
                    blocks.append(
                        ParsedBlock(
                            text="\n".join(shape_text),
                            slide_no=slide_index,
                            section_title=title,
                        )
                    )
        except Exception as exc:
            raise AppError("PARSE_FAILED", "PPTX parsing failed.") from exc
        return _ensure_non_empty(ParsedDocument(blocks=blocks, metadata={"parser": "pptx"}))


class TextParser(BaseParser):
    def parse(self, file_path: Path) -> ParsedDocument:
        try:
            raw_text = file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            raw_text = file_path.read_text(encoding="gb18030")
        except FileNotFoundError as exc:
            raise AppError("PARSE_FILE_NOT_FOUND", "Material file was not found.") from exc

        blocks: list[ParsedBlock] = []
        current_title: str | None = None
        paragraphs = re.split(r"\n\s*\n", raw_text)
        for paragraph in paragraphs:
            text = paragraph.strip()
            if not text:
                continue
            first_line = text.splitlines()[0].strip()
            if _looks_like_heading(first_line):
                current_title = first_line[:255]
            blocks.append(ParsedBlock(text=text, section_title=current_title))
        return _ensure_non_empty(ParsedDocument(blocks=blocks, metadata={"parser": "text"}))


class ImageParser(BaseParser):
    def parse(self, file_path: Path) -> ParsedDocument:
        raise AppError(
            "PARSE_UNSUPPORTED_TYPE",
            "Image OCR is reserved for a later extension.",
            details={"path": file_path.name},
        )


class ParserFactory:
    _parsers: dict[str, type[BaseParser]] = {
        "pdf": PdfParser,
        "docx": DocxParser,
        "pptx": PptxParser,
        "txt": TextParser,
        "md": TextParser,
        "image": ImageParser,
    }

    @classmethod
    def create(cls, file_type: str) -> BaseParser:
        parser_type = cls._parsers.get(file_type.lower())
        if parser_type is None:
            raise AppError(
                "PARSE_UNSUPPORTED_TYPE",
                "Unsupported material parser type.",
                details={"file_type": file_type},
            )
        return parser_type()


def parse_material_file(file_path: Path, file_type: str) -> ParsedDocument:
    return ParserFactory.create(file_type).parse(file_path)


def clean_text(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    normalized = re.sub(r"[\t\f\v]+", " ", normalized)
    normalized = re.sub(r" *\n *", "\n", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    normalized = re.sub(r"[ ]{2,}", " ", normalized)
    return normalized.strip()


def _looks_like_heading(text: str) -> bool:
    if not text or len(text) > 80:
        return False
    return bool(re.match(r"^(#{1,6}\s+|第.+[章节]|[0-9]+(\.[0-9]+)*\s+).+", text))


def _ensure_non_empty(document: ParsedDocument) -> ParsedDocument:
    document.blocks = [
        ParsedBlock(
            text=clean_text(block.text),
            page_no=block.page_no,
            slide_no=block.slide_no,
            section_title=block.section_title,
            metadata=block.metadata,
        )
        for block in document.blocks
        if clean_text(block.text)
    ]
    if not document.blocks:
        raise AppError("PARSE_EMPTY_CONTENT", "No readable text was found in material.")
    return document
